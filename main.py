from getter import *
import openpyxl

class Student:
    fi0 = ""
    group = ""
    id = 0
    ind = 0

    def __init__(self, fi0, gr, id, ind):
        self.fi0 = fi0
        self.group = gr
        self.id = id
        self.ind = ind

    def __str__(self):
        return self.fi0 + " " + self.group + " " + str(self.id)



# Open Excel file
wb = openpyxl.load_workbook('students.xlsx')
groups = ["Задачи"]
students = {}
for group in groups:
    sheet = wb[group]  # Access sheet by name
    # Read cell value
    ind = 2
    cell_value = sheet['B'+str(ind)].value
    #print(cell_value)
    while cell_value != None:
        id = 0
        if sheet['C' + str(ind)].value != None:
            id = int(sheet['C' + str(ind)].value)
            stud = Student(sheet['B' + str(ind)].value, group, id, ind)
            students[id] = stud
        ind += 1
        cell_value = sheet['B' + str(ind)].value
    print(ind)

wb.close()


lk_work = [
    {"link":"https://informatics.msk.ru/mod/statements/view.php?id=95930&standing", "table":"", "row":4, "name":"Lin"},
    {"link":"https://informatics.msk.ru/mod/statements/view.php?id=95931&standing", "table":"", "row":5, "name":"If"},
    {"link":"https://informatics.msk.ru/mod/statements/view.php?id=103033&standing", "table":"", "row":6, "name":"While"},
    {"link":"https://informatics.msk.ru/mod/statements/view.php?id=103034&standing", "table":"", "row":7, "name":"For"},
    {"link":"https://informatics.msk.ru/mod/statements/view.php?id=88939&standing", "table":"", "row":8, "name":"Array"},
    {"link":"https://informatics.msk.ru/mod/statements/view.php?id=96495&standing", "table":"", "row":9, "name":"Matrix"},
    {"link": "https://informatics.msk.ru/mod/statements/view.php?id=88099&standing", "table": "", "row": 10, "name": "Problems"}
]

student_data = get_table(lk_work)



wb_out = openpyxl.load_workbook('students.xlsx')
no_id_students = set()
for i in range(len(student_data)):
    job = student_data[i]
    max_tasks = job[1]
    col_res = lk_work[i]["row"]
    for stud in job[2]:
        stud_id = stud[0]
        if stud_id in students:
            student = students[stud_id]
            sheet = wb_out[student.group]
            sheet.cell(row=student.ind, column=col_res).value = round(stud[1]/max_tasks, 1)
        else:
            no_id_students.add(stud_id)
            #print("Error: No ID in table", stud_id)

with open("miss_id.txt","w") as file:
    for stud in no_id_students:
        file.write(str(stud)+"\n")

wb_out.save('students.xlsx')
wb_out.close()